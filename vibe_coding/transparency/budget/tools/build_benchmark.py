#!/usr/bin/env python3
""""מה מקבלים תמורת הכסף": השוואה של פרדס חנה-כרכור לרשויות דומות, לפי עלות ליחידת תוצאה.

שני מקורות פתוחים, לכל הרשויות בארץ:
- **קובצי הנתונים לעיבוד של הלמ"ס, "הרשויות המקומיות בישראל"** (src/benchmark/cbs/libud_*):
  אוכלוסייה, אשכול חברתי-כלכלי, פסולת (ק"ג ליום לנפש), אחוז מחזור, תלמידים, זכאות לבגרות,
  והוצאות החינוך בתקציב הרגיל. הקבצים יושבים ב-DocLib/2019/hamakomiot1999_2017/ באתר הלמ"ס,
  שחוסם הורדה שלא מדפדפן; הם ירדו בכרום (fetch מתוך דף של הלמ"ס).
- **הדוחות הכספיים המבוקרים של משרד הפנים** (data.gov.il, "local-authorities"): תשלומי פרק
  "תברואה" בטופס 2, כלומר כמה הרשות הוציאה בפועל על פינוי אשפה וניקיון. לא נמצא בקובץ הלמ"ס.

מה מחושב לכל רשות ושנה:
- פסולת: טון בשנה = ק"ג ליום לנפש x 365 x אוכלוסייה / 1000. שקל לטון = תשלומי תברואה / טון.
  שקל לתושב = תשלומי תברואה / אוכלוסייה. ולצידם אחוז המחזור.
- חינוך: שקל לתלמיד = הוצאות חינוך בתקציב הרגיל / תלמידים, ולצידו אחוז הזכאים לבגרות.

רשויות דומות: עיריות ומועצות מקומיות באשכול 6–8, עם 20–90 אלף תושבים. זו הבחירה הכי פשוטה
שאפשר להסביר בשורה אחת; הדף מציג את הרשימה.

    python3 tools/build_benchmark.py   ->   web/data/benchmark.json
"""
import glob
import json
import os
import re
import statistics

import openpyxl
import requests

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "src", "benchmark")
OUT = os.path.join(HERE, "..", "web", "data", "benchmark.json")
HOME = "פרדס חנה-כרכור"
UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) Chrome/128 phk-transparency"}

# דוחות מבוקרים, משאב לכל שנת דוח. כל דוח נותן גם את השנה הקודמת, ולכן 2019 מכסה את 2018,
# 2022 את 2021, ו-2020 חסרה (אין דוח 2020 או 2021 במאגר).
AUDITED = {"2024": "6e153ddd-d4d4-4450-a78c-20f594a7986a", "2023": "5825080f-905c-40f7-9f53-8ea1cf8ed394",
           "2022": "e5ff9ad0-6db2-4660-a94e-4499fce9475d", "2019": "f3f39a40-8404-4179-bcee-616ab2433631",
           "2018": "1618b871-9b3d-4bc5-acc0-c1fbe939a3b1"}

FIELDS = {  # שדה: ביטוי בכותרת המשורשרת של העמודה, בגיליון הפיזי או בגיליון התקציב
    "pop": (r"^סה\"כ ?אוכלוסייה בסוף השנה", "phys"),
    "cluster": (r"חברתי.?כלכלי.*אשכול \(מ-1 עד 10, 1 הנמוך", "phys"),
    "waste": (r"פסולת.*ליום לנפש", "phys"),
    "recycle": (r"אחוז מחזור", "phys"),
    "pupils": (r"\| תלמידים \| סה\"כ$", "phys"),
    "bagrut": (r"אחוז זכאים לתעודת בגרות מבין תלמידי", "phys"),
    "status": (r"^מעמד מוניציפלי", "phys"),
    "edu_pay": (r"שירותים ממלכתיים: חינוך", "budget"),
}


def norm(name):
    name = re.sub(r"[\s\-–־'\"״׳]+", "", str(name or ""))
    return name.replace("קריית", "קרית").replace("יי", "י")


def is_code(v):
    return isinstance(v, (int, float)) or (isinstance(v, str) and v.strip().isdigit())


def headers(rows):
    """הכותרת של כל עמודה, משורשרת מכל שורות הכותרת (עד השורה הראשונה עם סמל רשות)."""
    depth = next((i for i, r in enumerate(rows) if len(r) > 1 and r[0] and is_code(r[1])), 6)
    depth = min(depth, next((i for i, r in enumerate(rows[3:], 3) if r and str(r[0] or "").startswith("כלל ארצי")), depth))
    ncol = max(len(r) for r in rows[:12])
    return [" | ".join(re.sub(r"\s+", " ", str(rows[i][c])).strip() for i in range(depth)
                       if c < len(rows[i]) and rows[i][c] not in (None, "")) for c in range(ncol)]


def num(v):
    try:
        f = float(v)
        return f if f == f else None
    except (TypeError, ValueError):
        return None


def read_cbs(path):
    """{שם מנורמל: {שדה: ערך}} מקובץ הלמ"ס של שנה אחת."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = {"phys": next(n for n in wb.sheetnames if n.startswith("נתונים פיזיים")),
              "budget": next((n for n in wb.sheetnames if n.startswith("נתוני תקציב")), None)}
    out = {}
    for kind, sn in sheets.items():
        if not sn:
            continue
        rows = [list(r) for r in wb[sn].iter_rows(values_only=True)]
        hdr = headers(rows)
        cols = {}
        for f, (rx, where) in FIELDS.items():
            if where == kind:
                hit = [c for c, h in enumerate(hdr) if re.search(rx, h)]
                if hit:
                    cols[f] = hit[0]
        for r in rows[4:]:
            if not r or len(r) < 2 or not r[0] or not is_code(r[1]):
                continue  # שורות סיכום ("כלל ארצי", "עיריות סך הכל") בלי סמל רשות
            rec = out.setdefault(norm(r[0]), {"name": str(r[0]).strip(), "code": int(r[1])})
            for f, c in cols.items():
                v = r[c] if c < len(r) else None
                rec[f] = v.strip() if f == "status" and isinstance(v, str) else num(v)
    return out


def audited_sanitation():
    """{שנה: {שם מנורמל: אלפי ש"ח}}: תשלומי תברואה בפועל, מטופס 2 של הדוח המבוקר."""
    out = {}
    for rep, rid in AUDITED.items():
        off = 0
        while True:
            d = requests.get("https://data.gov.il/api/3/action/datastore_search", headers=UA, timeout=120,
                             params={"resource_id": rid, "limit": 5000, "offset": off,
                                     "filters": json.dumps({"שורה": "תברואה"}, ensure_ascii=False)}).json()["result"]
            for r in d["records"]:
                col, sheet = r.get("עמודה", ""), r.get("גליון") or r.get("גיליון")
                if sheet != "טופס 2" or not col.startswith("תשלומים - ביצוע"):
                    continue
                year = str(r.get("שנת_נתונים"))
                v = num(r.get("ערך"))
                if v is not None:
                    out.setdefault(year, {}).setdefault(norm(r["שם_רשות"]), v)
            off += len(d["records"])
            if off >= d["total"] or not d["records"]:
                break
    return out


def main():
    cbs = {}
    for path in glob.glob(os.path.join(SRC, "cbs", "libud_*.xlsx")):
        m = re.search(r"(\d{2,4})\.xlsx$", path)
        y = int(m.group(1))
        y = y + 2000 if y < 100 else y
        if y >= 2016:
            cbs[y] = read_cbs(path)
            print(y, len(cbs[y]), "רשויות", flush=True)
    san = audited_sanitation()
    print("תברואה:", {y: len(v) for y, v in sorted(san.items())}, flush=True)

    years = sorted(cbs)
    rows = []
    for y in years:
        for k, r in cbs[y].items():
            if r.get("status") not in ("עירייה", "מועצה מקומית") or not r.get("pop"):
                continue
            pop = r["pop"] * (1000 if r["pop"] < 1000 else 1)  # 2016–2017: באלפים
            rec = {"y": y, "n": r["name"], "pop": round(pop), "cl": r.get("cluster")}
            s = san.get(str(y), {}).get(k)
            if r.get("waste") and s:
                tons = r["waste"] * 365 * pop / 1000
                rec.update(kg=round(r["waste"], 2), tons=round(tons), san=round(s),
                           per_ton=round(s * 1000 / tons), per_res=round(s * 1000 / pop))
            if r.get("recycle") is not None:
                rec["rec"] = round(r["recycle"], 1)
            if r.get("edu_pay") and r.get("pupils"):
                rec.update(edu=round(r["edu_pay"]), pupils=round(r["pupils"]),
                           per_pupil=round(r["edu_pay"] * 1000 / r["pupils"]))
            if r.get("bagrut") is not None:
                rec["bag"] = round(r["bagrut"], 1)
            rows.append(rec)

    # האשכול נמדד פעם בכמה שנים; רשות בלי אשכול בשנה מסוימת מקבלת את האחרון הידוע שלה
    last = {}
    for r in sorted(rows, key=lambda r: r["y"]):
        if r.get("cl"):
            last[r["n"]] = r["cl"]
    for r in rows:
        r["cl"] = r.get("cl") or last.get(r["n"])
    peers = sorted({r["n"] for r in rows if r["y"] == years[-1] and r.get("cl") and 6 <= r["cl"] <= 8
                    and 20000 <= r["pop"] <= 90000 and r["n"] != HOME})

    def med(y, key):
        vals = [r[key] for r in rows if r["y"] == y and r["n"] in peers and r.get(key) is not None]
        if len(vals) < 4:
            return None
        q = statistics.quantiles(vals, n=4)
        return {"q1": round(q[0], 1), "med": round(statistics.median(vals), 1), "q3": round(q[2], 1), "n": len(vals)}

    summary = {y: {k: med(y, k) for k in ("per_ton", "per_res", "kg", "rec", "per_pupil", "bag")} for y in years}
    json.dump({"home": HOME, "peers": peers, "peer_rule": "עיריות ומועצות מקומיות באשכול חברתי-כלכלי 6–8, עם 20–90 אלף תושבים",
               "years": years, "summary": summary, "rows": rows,
               "sources": {"cbs": "https://www.cbs.gov.il/he/publications/Pages/2019/הרשויות-המקומיות-בישראל-קובצי-נתונים-לעיבוד-1999-2017.aspx",
                           "audited": "https://data.gov.il/dataset/local-authorities"}},
              open(OUT, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    home = [r for r in rows if r["n"] == HOME]
    for r in home:
        print(r)
    print(len(peers), "רשויות דומות:", ", ".join(peers))


if __name__ == "__main__":
    main()
