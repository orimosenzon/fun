#!/usr/bin/env python3
"""Download a Google Sheet (shared "anyone with the link") as xlsx and export every tab's
cached values + background colors to JSON, for the local harness (tools/harness.js).

    python3 tools/export_sheet.py <spreadsheet_id> [out_dir]

Output goes to out_dir (default tools/out/, git-ignored): <id>.xlsx and sheets.json.
The export contains the PA's real course data, so it must stay out of git.
"""
import datetime
import json
import pathlib
import subprocess
import sys

import openpyxl

MAX_ROWS = 60
MAX_COLS = 26


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    sheet_id = sys.argv[1]
    out_dir = pathlib.Path(sys.argv[2] if len(sys.argv) > 2 else pathlib.Path(__file__).parent / "out")
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / f"{sheet_id}.xlsx"
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    subprocess.run(["curl", "-sL", url, "-o", str(xlsx)], check=True)

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows, bgs = [], []
        for r in range(1, min(ws.max_row, MAX_ROWS) + 1):
            vals, cols = [], []
            for c in range(1, MAX_COLS + 1):
                cell = ws.cell(r, c)
                v = cell.value
                if isinstance(v, datetime.datetime):
                    v = {"$date": v.strftime("%Y-%m-%d")}
                elif v is None:
                    v = ""
                vals.append(v)
                fill = cell.fill
                solid = fill is not None and fill.fill_type == "solid" and isinstance(fill.fgColor.rgb, str)
                rgb = fill.fgColor.rgb if solid else "FFFFFFFF"
                cols.append("#" + rgb[-6:].lower())
            rows.append(vals)
            bgs.append(cols)
        out[ws.title] = {"values": rows, "backgrounds": bgs}
    (out_dir / "sheets.json").write_text(json.dumps(out))
    print(f"{len(out)} tabs -> {out_dir / 'sheets.json'}")


if __name__ == "__main__":
    main()
